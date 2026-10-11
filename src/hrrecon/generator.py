"""Generador de datos FICTICIOS para demostración y pruebas.

Produce una nómina (Excel) y una exportación de directorio (CSV con ';') de una empresa
inventada, "Comercial Andina (ficticia)", con dominio reservado ``.example``. Inyecta casos
trampa a propósito y guarda en ``expected_traps.json`` exactamente qué debe detectar cada
regla, para que las pruebas lo verifiquen.

Ningún dato corresponde a personas reales: nombres y apellidos se combinan al azar y los
códigos de empleado son internos de 6 dígitos (no son cédulas).
"""

from __future__ import annotations

import csv
import json
import random
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .normalize import strip_accents

DOMAIN = "comercialandina.example"
DN_SUFFIX = "OU=Usuarios,DC=comercialandina,DC=example"

# fmt: off
FIRST_NAMES_M = [
    "Andrés", "José", "Sebastián", "Martín", "Nicolás", "Matías", "Julián", "Agustín", "Ramón",
    "Germán", "Iván", "Rubén", "Jesús", "Efraín", "Hernán", "Darío", "Simón", "Cristóbal",
    "Fabián", "Joaquín", "Ángel", "Esteban", "Patricio", "Santiago", "Gonzalo", "Mauricio",
    "Wilson", "Byron", "Édison", "Rodrigo", "Francisco", "Diego", "Jhonatan", "Óscar", "Raúl",
]
FIRST_NAMES_F = [
    "María", "Belén", "Inés", "Mónica", "Verónica", "Lucía", "Sofía", "Ximena", "Gissela",
    "Mishell", "Dayana", "Jhoana", "Paola", "Carolina", "Gabriela", "Andrea", "Daniela", "Ruth",
    "Fernanda", "Estefanía", "Rocío", "Nathaly", "Karina", "Lorena", "Pamela", "Tatiana",
    "Alejandra", "Valeria", "Cristina", "Marisol", "Begoña", "Noemí", "Johanna", "Jimena",
]
SURNAMES = [
    "Muñoz", "Peñaherrera", "Ordóñez", "Núñez", "Yépez", "Ibarra", "Chávez", "Guzmán",
    "Cevallos", "Paredes", "Zambrano", "Montúfar", "Albán", "Bermúdez", "Andrade", "Aguirre",
    "Benítez", "Cárdenas", "Espín", "Jácome", "Lozada", "Maldonado", "Narváez", "Pazmiño",
    "Quiñónez", "Rodríguez", "Salazar", "Tapia", "Villacís", "Zúñiga", "Proaño", "Galárraga",
    "Endara", "Toapanta", "Chiriboga", "Quishpe", "Guamán", "Tituaña", "Cañar", "Llerena",
    "Viteri", "Íñiguez", "Sarmiento", "Mejía", "Echeverría", "Ponce", "Avilés", "Córdova",
    "Herrera", "Vásconez", "Ayala", "Calderón", "Pérez", "Logroño", "Robalino", "Garcés",
]
CITIES = ["Quito", "Guayaquil", "Cuenca", "Ambato", "Manta", "Loja"]

# departamento: (cargo del jefe, [(cargo, peso)], cargo de supervisor o None)
DEPARTMENTS: dict[str, tuple[str, list[tuple[str, int]], str | None]] = {
    "Talento Humano": (
        "Jefe de Talento Humano",
        [("Analista de Nómina", 3), ("Asistente de Talento Humano", 2),
         ("Psicóloga Organizacional", 1)],
        None,
    ),
    "Finanzas": (
        "Jefe Financiero",
        [("Contador", 2), ("Analista Contable", 4), ("Asistente Contable", 3), ("Tesorero", 1)],
        None,
    ),
    "Tecnología": (
        "Jefe de Tecnología",
        [("Desarrollador", 4), ("Analista de Sistemas", 3), ("Técnico de Soporte", 4),
         ("Administrador de Redes", 1)],
        None,
    ),
    "Comercial": (
        "Gerente Comercial",
        [("Ejecutivo de Ventas", 6), ("Asesor Comercial", 6), ("Analista de Precios", 1)],
        "Supervisor de Ventas",
    ),
    "Operaciones": (
        "Jefe de Operaciones",
        [("Auxiliar de Bodega", 8), ("Chofer", 4), ("Montacarguista", 2),
         ("Analista de Logística", 2)],
        "Supervisor de Bodega",
    ),
    "Atención al Cliente": (
        "Jefe de Atención al Cliente",
        [("Agente de Atención al Cliente", 8), ("Analista de Reclamos", 2)],
        "Coordinador de Call Center",
    ),
    "Marketing": (
        "Jefe de Marketing",
        [("Diseñador Gráfico", 2), ("Analista de Marketing Digital", 2), ("Community Manager", 1)],
        None,
    ),
    "Adquisiciones": (
        "Jefe de Adquisiciones",
        [("Analista de Compras", 3), ("Asistente de Compras", 2)],
        None,
    ),
    "Asesoría Jurídica": (
        "Jefe de Asesoría Jurídica",
        [("Abogado", 2), ("Asistente Legal", 1)],
        None,
    ),
    "Auditoría Interna": (
        "Auditor Interno Jefe",
        [("Auditor", 2), ("Analista de Riesgos", 1)],
        None,
    ),
}
DEPT_WEIGHTS = {
    "Talento Humano": 4, "Finanzas": 6, "Tecnología": 7, "Comercial": 14, "Operaciones": 14,
    "Atención al Cliente": 10, "Marketing": 3, "Adquisiciones": 3, "Asesoría Jurídica": 2,
    "Auditoría Interna": 2,
}

SERVICE_ACCOUNTS = [
    ("svc_backup", "Cuenta de servicio - respaldos"),
    ("svc-sql01", "Cuenta de servicio - SQL Server"),
    ("svc_nomina", "Cuenta de servicio - integración nómina"),
    ("sa_scanner", "Cuenta de servicio - escáneres"),
    ("admin.ti", "Administrador de dominio"),
    ("admin.redes", "Administrador de red"),
    ("scanner.piso3", "Escáner multifunción piso 3"),
    ("backup.agent", "Agente de respaldo"),
]
SPECIAL_ACCOUNTS = [
    ("recepcion.general", "Recepción General", "Cuenta compartida de recepción"),
    ("sala.reuniones", "Sala de Reuniones", "Buzón de recurso"),
]
ORPHANS = [  # (cuenta, nombre, employeeID, habilitada)
    ("contratista.ext01", "Contratista Externo 01", "", True),
    ("ex.consultor", "Consultor Externo", "", True),
    ("pruebas.qa", "Usuario de Pruebas QA", "", True),
    ("wvaca", "Washington Vaca Ruales", "099901", True),
    ("lsimbana", "Lorena Simbaña Toro", "099902", True),
    ("temporal.bodega", "Temporal Bodega", "", False),
    ("ex.pasante2023", "Ex Pasante 2023", "", False),
]
# fmt: on


@dataclass
class Person:
    emp_id: int
    first: str
    last: str
    gender: str
    department: str
    title: str
    manager_id: int | None
    city: str
    hire: date
    status: str = "Activo"
    term: date | None = None
    email: str = ""
    account: str | None = None
    # Valores del directorio (por defecto iguales a la nómina)
    enabled: bool = True
    dir_title: str | None = None
    dir_department: str | None = None
    dir_manager_id: int | None = None
    dir_email: str | None = None
    dir_emp_id: str | None = None  # None = usar el código real
    last_logon: datetime | None = None
    hr_title: str | None = None  # para ruido cosmético / vacíos en nómina
    hr_department: str | None = None
    hr_email: str | None = None
    hr_status: str | None = None
    hr_id_override: str | None = None
    extra_accounts: list[str] = field(default_factory=list)
    tags: list[str] = field(default_factory=list)

    @property
    def display_name(self) -> str:
        return f"{self.first.split()[0]} {self.last}"

    @property
    def code(self) -> str:
        return f"{self.emp_id:06d}"


def _ascii(text: str) -> str:
    return strip_accents(text).lower().replace(" ", "")


class _Builder:
    def __init__(self, n: int, seed: int, as_of: date):
        self.rng = random.Random(seed)
        self.n = n
        self.as_of = as_of
        self.people: list[Person] = []
        self.extra_hr_rows: list[Person] = []
        self.used_names: set[str] = set()
        self.used_accounts: set[str] = {a for a, _ in SERVICE_ACCOUNTS}
        self.used_accounts |= {a for a, _, _ in SPECIAL_ACCOUNTS} | {o[0] for o in ORPHANS}
        self.used_emails: set[str] = set()
        self.used_ids: set[int] = set()
        self.expected: dict = {}

    # ---------- helpers ----------
    def days_ago(self, lo: int, hi: int) -> date:
        return self.as_of - timedelta(days=self.rng.randint(lo, hi))

    def logon_on(self, day: date) -> datetime:
        return datetime(
            day.year,
            day.month,
            day.day,
            self.rng.randint(7, 19),
            self.rng.randint(0, 59),
            self.rng.randint(0, 59),
        )

    def new_id(self) -> int:
        while True:
            value = self.rng.randint(1000, 98999)
            if value not in self.used_ids:
                self.used_ids.add(value)
                return value

    def new_name(self) -> tuple[str, str, str]:
        while True:
            gender = self.rng.choice("MF")
            pool = FIRST_NAMES_M if gender == "M" else FIRST_NAMES_F
            first = " ".join(self.rng.sample(pool, 2))
            last = " ".join(self.rng.sample(SURNAMES, 2))
            display = f"{first.split()[0]} {last}"
            if display not in self.used_names:
                self.used_names.add(display)
                return first, last, gender

    def make_account(self, p: Person) -> str:
        first = _ascii(p.first.split()[0])
        s1, s2 = (_ascii(s) for s in p.last.split())
        for candidate in (f"{first[0]}{s1}", f"{first[0]}{s1}{s2[0]}", f"{first}.{s1}"):
            if candidate not in self.used_accounts:
                self.used_accounts.add(candidate)
                return candidate
        i = 2
        while f"{first[0]}{s1}{i}" in self.used_accounts:
            i += 1
        self.used_accounts.add(f"{first[0]}{s1}{i}")
        return f"{first[0]}{s1}{i}"

    def make_email(self, p: Person) -> str:
        first = _ascii(p.first.split()[0])
        s1, s2 = (_ascii(s) for s in p.last.split())
        for local in (f"{first}.{s1}", f"{first}.{s1}{s2[0]}", f"{first[0]}{s1}{s2}"):
            email = f"{local}@{DOMAIN}"
            if email not in self.used_emails:
                self.used_emails.add(email)
                return email
        raise RuntimeError("no se pudo generar correo único")

    def person(self, department: str, title: str, manager_id: int | None) -> Person:
        first, last, gender = self.new_name()
        p = Person(
            emp_id=self.new_id(),
            first=first,
            last=last,
            gender=gender,
            department=department,
            title=title,
            manager_id=manager_id,
            city=self.rng.choice(CITIES),
            hire=self.days_ago(200, 4000),
        )
        p.email = self.make_email(p)
        p.account = self.make_account(p)
        p.last_logon = self.logon_on(self.days_ago(0, 30))
        self.people.append(p)
        return p

    # ---------- organización base ----------
    def build_org(self) -> None:
        gm = self.person("Gerencia General", "Gerente General", None)
        gm.hire = self.days_ago(3000, 5000)
        self.heads: dict[str, Person] = {}
        self.supervisors: dict[str, Person] = {}
        for dept, (head_title, _, sup_title) in DEPARTMENTS.items():
            self.heads[dept] = self.person(dept, head_title, gm.emp_id)
            if sup_title:
                self.supervisors[dept] = self.person(dept, sup_title, self.heads[dept].emp_id)
        self.managers = [gm, *self.heads.values(), *self.supervisors.values()]

        # Pasantes: 10, reportan al jefe del departamento.
        self.interns: list[Person] = []
        for _ in range(10):
            dept = self.rng.choice(list(DEPARTMENTS))
            p = self.person(dept, "Pasante", self.heads[dept].emp_id)
            p.hire = self.days_ago(35, 150)
            self.interns.append(p)

        depts = list(DEPT_WEIGHTS)
        weights = [DEPT_WEIGHTS[d] for d in depts]
        self.staff: list[Person] = []
        while len(self.people) < self.n:
            dept = self.rng.choices(depts, weights)[0]
            titles, tw = zip(*DEPARTMENTS[dept][1], strict=True)
            title = self.rng.choices(titles, tw)[0]
            boss = self.supervisors.get(dept) if self.rng.random() < 0.6 else None
            boss = boss or self.heads[dept]
            self.staff.append(self.person(dept, title, boss.emp_id))
        self.pool = list(self.staff)
        self.rng.shuffle(self.pool)

    def take(self, k: int) -> list[Person]:
        taken, self.pool = self.pool[:k], self.pool[k:]
        return taken

    # ---------- casos trampa ----------
    def inject(self) -> None:
        rng, exp = self.rng, self.expected
        a = self.as_of

        def terminate(p: Person, days: int | None) -> None:
            p.status = "Baja"
            p.term = a - timedelta(days=days) if days is not None else None
            if p.term:
                p.hire = min(p.hire, p.term - timedelta(days=200))
                p.last_logon = self.logon_on(p.term - timedelta(days=rng.randint(1, 15)))

        # R1: bajas con cuenta habilitada
        r1_crit, r1_high = [], []
        for p in self.take(8):
            terminate(p, rng.randint(10, 200))
            r1_crit.append(p)
        for p, d in zip(self.take(3), (7, 2, 4), strict=True):  # 7 = umbral exacto -> ALTA
            terminate(p, d)
            r1_high.append(p)
        (nodate,) = self.take(1)
        terminate(nodate, None)
        nodate.last_logon = self.logon_on(self.days_ago(20, 40))
        r1_high.append(nodate)
        (incons,) = self.take(1)  # Estado "Activo" pero fecha de salida pasada
        incons.term = a - timedelta(days=30)
        incons.last_logon = self.logon_on(a - timedelta(days=35))
        r1_crit.append(incons)

        # R6: inicio de sesión posterior a la baja
        r6_enabled = self.take(3)
        for p in r6_enabled:
            terminate(p, rng.randint(20, 60))
            p.last_logon = self.logon_on(p.term + timedelta(days=rng.randint(2, 10)))
            r1_crit.append(p)
        r6_disabled = self.take(2)
        for p in r6_disabled:
            terminate(p, 40)
            p.enabled = False
            p.last_logon = self.logon_on(p.term + timedelta(days=3))

        # Bajas programadas (fecha futura) con cuenta habilitada -> no se reportan
        scheduled = self.take(3)
        for p in scheduled:
            p.status = "Baja"
            p.term = a + timedelta(days=rng.randint(5, 30))

        # Bajas correctas: cuenta deshabilitada o sin cuenta -> no se reportan
        for p in self.take(25):
            terminate(p, rng.randint(10, 900))
            p.enabled = False
        for p in self.take(5):
            terminate(p, rng.randint(400, 1500))
            p.account = None

        # R2 / R2b
        r2 = self.take(8)
        for p in r2:
            p.account = None
            p.hire = self.days_ago(15, 120)
        future_hires = []
        for _ in range(2):
            dept = rng.choice(list(DEPARTMENTS))
            p = self.person(dept, DEPARTMENTS[dept][1][0][0], self.heads[dept].emp_id)
            p.hire = a + timedelta(days=10)
            p.account = None
            future_hires.append(p)
        r2b = self.take(4)
        for p in r2b:
            p.enabled = False

        interns_no_account = self.interns[:6]
        for p in interns_no_account:
            p.account = None
        self.interns[0].hr_title = "PASANTE "
        self.interns[1].hr_title = "pasante"

        # R3: diferencias de atributos
        drift: dict[str, list[Person]] = {"title": [], "department": [], "manager": [], "email": []}
        for p in self.take(8):
            others = [t for t, _ in DEPARTMENTS[p.department][1] if t != p.title]
            p.dir_title = rng.choice(others)  # ascenso/cambio no reflejado en el directorio
            drift["title"].append(p)
        for p in self.take(6):
            p.dir_department = rng.choice([d for d in DEPARTMENTS if d != p.department])
            drift["department"].append(p)
        for p in self.take(6):
            others = [m for m in self.managers if m.emp_id not in (p.manager_id, p.emp_id)]
            p.dir_manager_id = rng.choice(others).emp_id
            drift["manager"].append(p)
        for p in self.take(4):
            first = _ascii(p.first.split()[0])
            s1, s2 = (_ascii(s) for s in p.last.split())
            p.dir_email = f"{first}.{s2}{s1}@{DOMAIN}"  # correo anterior / con apellidos invertidos
            drift["email"].append(p)
        empty_dept = self.take(3)
        for p in empty_dept:
            p.dir_department = ""
            drift["department"].append(p)
        empty_title = self.take(2)
        for p in empty_title:
            p.hr_title = ""
            drift["title"].append(p)

        # Emparejamiento por correo: employeeID vacío en el directorio
        by_email = self.take(3)
        for p in by_email:
            p.dir_emp_id = ""

        # R5: calidad de datos
        rehire, dup_a = self.take(2)
        old = Person(**{**rehire.__dict__, "extra_accounts": [], "tags": []})
        old.status, old.hire = "Baja", rehire.hire - timedelta(days=1500)
        old.term, old.account = rehire.hire - timedelta(days=400), None
        self.extra_hr_rows.append(old)
        first, last, gender = self.new_name()
        dup_b = Person(
            emp_id=dup_a.emp_id,
            first=first,
            last=last,
            gender=gender,
            department=dup_a.department,
            title=dup_a.title,
            manager_id=dup_a.manager_id,
            city=dup_a.city,
            hire=self.days_ago(60, 300),
            account=None,
        )
        dup_b.email = self.make_email(dup_b)
        self.extra_hr_rows.append(dup_b)

        multi = self.take(2)
        for p in multi:
            name, i = f"{p.account}.2", 2
            while name in self.used_accounts:
                i += 1
                name = f"{p.account}.{i}"
            self.used_accounts.add(name)
            p.extra_accounts.append(name)

        dup_email_hr = []
        for _ in range(2):
            p1, p2 = self.take(2)
            p2.hr_email = p1.email  # copia/pega en nómina
            dup_email_hr.append((p1, p2))
            drift["email"].append(p2)
        q1, q2 = self.take(2)
        q2.dir_email = q1.email  # dos cuentas con el mismo correo en el directorio
        drift["email"].append(q2)

        (empty_id,) = self.take(1)
        empty_id.hr_id_override = ""
        empty_id.account = None
        (unknown_status,) = self.take(1)
        unknown_status.hr_status = ""

        # R7: cuentas habilitadas sin uso
        stale = self.take(8)
        for p in stale:
            p.last_logon = self.logon_on(self.days_ago(100, 400))
        never = self.take(2)
        for p in never:
            p.last_logon = None
            p.hire = self.days_ago(300, 900)
        recent = self.take(3)  # altas recientes sin inicio de sesión: NO se reportan
        for p in recent:
            p.last_logon = None
            p.hire = self.days_ago(1, 20)

        # Ruido cosmético que NO debe generar hallazgos
        noise = self.take(60)
        variants = ("title_upper", "dept_spaces", "email_upper", "title_spaces", "dept_lower")
        for i, p in enumerate(noise):
            kind = variants[i % len(variants)]
            if kind == "title_upper":
                p.dir_title = strip_accents(p.title).upper()
            elif kind == "dept_spaces":
                p.hr_department = f"  {p.department} "
            elif kind == "email_upper":
                p.dir_email = p.email.upper()
            elif kind == "title_spaces":
                p.hr_title = p.title.replace(" ", "  ", 1) + " "
            else:
                p.dir_department = strip_accents(p.department).lower()

        def ids(people):
            return sorted(p.code for p in people)

        exp.update(
            {
                "R1": {"critical": ids(r1_crit), "high": ids(r1_high)},
                "R2": ids(r2),
                "R2b": ids(r2b),
                "R3": {k: ids(v) for k, v in drift.items()},
                "R4": sorted(o[0] for o in ORPHANS),
                "R4_enabled": sorted(o[0] for o in ORPHANS if o[3]),
                "R5": {
                    "duplicate_hr_id": ids([rehire, dup_a]),
                    "multiple_enabled_accounts": ids(multi),
                    "duplicate_email_hr": sorted(p1.email for p1, _ in dup_email_hr),
                    "duplicate_email_directory": [q1.email],
                    "missing_employee_id": 1,
                    "unknown_status": ids([unknown_status]),
                    "matched_by_email": ids(by_email),
                    "terminated_without_date": ids([nodate]),
                    "status_date_inconsistent": ids([incons]),
                },
                "R6": sorted(p.account for p in r6_enabled + r6_disabled),
                "R7": sorted(p.account for p in stale + never),
                "excluded": {
                    "service_accounts": sorted(s for s, _ in SERVICE_ACCOUNTS),
                    "special_accounts": sorted(s for s, _, _ in SPECIAL_ACCOUNTS),
                    "no_account_required": ids(interns_no_account),
                    "scheduled_termination": ids(scheduled),
                    "scheduled_hire": ids(future_hires),
                },
                "no_findings": ids(noise + recent),
            }
        )


# ---------------- escritura de archivos ----------------


def _id_cell(rng: random.Random, code: str):
    """Formatos de ID inconsistentes, como en exportaciones reales."""
    if code == "":
        return ""
    r = rng.random()
    if r < 0.2:
        return int(code)  # Excel lo guarda como número: se pierden los ceros
    if r < 0.3:
        return f" {int(code)} "
    return code


def _date_cell(rng: random.Random, value: date | None):
    if value is None:
        return None
    return value if rng.random() < 0.5 else value.strftime("%d/%m/%Y")


def _status_cell(rng: random.Random, status: str) -> str:
    variants = {
        "Activo": ["Activo", "ACTIVO", " activo", "Activo"],
        "Baja": ["Baja", "BAJA", "Inactivo"],
    }
    return rng.choice(variants[status])


def _write_hr(b: _Builder, path: Path) -> int:
    rng = random.Random(b.rng.random())
    wb = Workbook()
    ws = wb.active
    ws.title = "Empleados"
    headers = [
        "Código Empleado",
        "Nombres",
        "Apellidos",
        "Estado",
        "Cargo ",
        "Departamento",
        "Código Jefe",
        "Correo Corporativo",
        "Fecha Ingreso",
        "Fecha Salida",
        "Ciudad",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    rows = sorted(b.people + b.extra_hr_rows, key=lambda p: (p.emp_id, p.status != "Baja"))
    for p in rows:
        code = p.code if p.hr_id_override is None else p.hr_id_override
        ws.append(
            [
                _id_cell(rng, code),
                p.first,
                p.last,
                p.hr_status if p.hr_status is not None else _status_cell(rng, p.status),
                p.hr_title if p.hr_title is not None else p.title,
                p.hr_department if p.hr_department is not None else p.department,
                _id_cell(rng, f"{p.manager_id:06d}") if p.manager_id else None,
                p.hr_email if p.hr_email is not None else p.email,
                _date_cell(rng, p.hire),
                _date_cell(rng, p.term),
                p.city,
            ]
        )
    for col, width in zip("ABCDEFGHIJK", (16, 24, 24, 10, 32, 22, 13, 40, 14, 14, 12), strict=True):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for idx in (8, 9):
            if isinstance(row[idx].value, date):
                row[idx].number_format = "DD/MM/YYYY"
    wb.save(path)
    return len(rows)


def _write_directory(b: _Builder, path: Path) -> int:
    rng = random.Random(b.rng.random())
    by_id = {p.emp_id: p for p in b.people}
    headers = [
        "sAMAccountName",
        "employeeID",
        "Enabled",
        "DisplayName",
        "GivenName",
        "Surname",
        "Title",
        "Department",
        "Manager",
        "mail",
        "LastLogonDate",
        "Description",
    ]

    def enabled_cell(flag: bool) -> str:
        return rng.choice(["True", "True", "TRUE"]) if flag else rng.choice(["False", "FALSE"])

    def logon_cell(value: datetime | None) -> str:
        if value is None:
            return ""
        fmt = "%Y-%m-%d %H:%M:%S" if rng.random() < 0.7 else "%d/%m/%Y %H:%M"
        return value.strftime(fmt)

    def manager_cell(manager_id: int | None) -> str:
        if not manager_id:
            return ""
        m = by_id[manager_id]
        r = rng.random()
        if r < 0.6:
            ou = strip_accents(m.department)
            return f"CN={m.display_name},OU={ou},{DN_SUFFIX}"
        if r < 0.85:
            return m.account or ""
        return m.code

    rows: list[list[str]] = []
    for p in b.people:
        if not p.account:
            continue
        emp_cell = (
            p.dir_emp_id
            if p.dir_emp_id is not None
            else (str(p.emp_id) if rng.random() < 0.5 else p.code)
        )
        base = [
            emp_cell,
            enabled_cell(p.enabled),
            p.display_name,
            p.first.split()[0],
            p.last,
            p.dir_title if p.dir_title is not None else p.title,
            p.dir_department if p.dir_department is not None else p.department,
            manager_cell(p.dir_manager_id or p.manager_id),
            p.dir_email if p.dir_email is not None else p.email,
            logon_cell(p.last_logon),
            "",
        ]
        rows.append([p.account, *base])
        for extra in p.extra_accounts:
            second = list(base)
            second[9] = logon_cell(b.logon_on(b.days_ago(0, 20)))
            rows.append([extra, *second])
    for name, desc in SERVICE_ACCOUNTS:
        logon = b.logon_on(b.days_ago(0, 300))
        rows.append(
            [
                name,
                "",
                enabled_cell(True),
                name,
                "",
                "",
                "",
                "Tecnología",
                "",
                "",
                logon_cell(logon),
                desc,
            ]
        )
    for name, display, desc in SPECIAL_ACCOUNTS:
        logon = b.logon_on(b.days_ago(150, 300))
        rows.append(
            [name, "", enabled_cell(True), display, "", "", "", "", "", "", logon_cell(logon), desc]
        )
    for name, display, emp, enabled in ORPHANS:
        logon = b.logon_on(b.days_ago(0, 30))
        rows.append(
            [
                name,
                emp,
                enabled_cell(enabled),
                display,
                "",
                "",
                "Consultor",
                "",
                "",
                f"{name}@{DOMAIN}",
                logon_cell(logon),
                "",
            ]
        )
    rng.shuffle(rows)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)
    return len(rows)


def generate(
    out_dir: str | Path, n_employees: int = 500, seed: int = 42, as_of: date | None = None
) -> dict:
    """Genera nomina.xlsx, directorio.csv y expected_traps.json en ``out_dir``."""
    if n_employees < 250:
        raise ValueError("Se necesitan al menos 250 empleados para inyectar todos los casos trampa")
    as_of = as_of or date.today()
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    b = _Builder(n_employees, seed, as_of)
    b.build_org()
    b.inject()
    hr_rows = _write_hr(b, out / "nomina.xlsx")
    dir_rows = _write_directory(b, out / "directorio.csv")
    expected = {
        "_note": "Datos ficticios. Qué debe detectar cada regla con la configuración de ejemplo.",
        "as_of": as_of.isoformat(),
        "seed": seed,
        "hr_rows": hr_rows,
        "directory_rows": dir_rows,
        **b.expected,
    }
    (out / "expected_traps.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return expected
