"""Libro Excel: hoja de resumen + una hoja por categoría + exclusiones + entradas."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import __version__
from ..models import Result
from .common import COLUMNS, SEVERITIES, SHEET_NAMES, categories, cell_value, key_messages

HEADER_FILL = PatternFill("solid", fgColor="1F4E8C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEVERITY_FILL = {
    "CRÍTICA": PatternFill("solid", fgColor="F6D3D3"),
    "ALTA": PatternFill("solid", fgColor="FBE0D5"),
    "MEDIA": PatternFill("solid", fgColor="FEF0CF"),
    "BAJA": PatternFill("solid", fgColor="ECECEA"),
}
MAX_WIDTH = 60
WRAP_COLUMNS = {"Detalle", "Motivo", "Nombre", "Empleado"}


def _table(ws: Worksheet, start_row: int, headers: list[str], rows: list[list]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, values in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            wrap = headers[c - 1] in WRAP_COLUMNS
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            if headers[c - 1] == "Severidad" and value in SEVERITY_FILL:
                cell.fill = SEVERITY_FILL[value]
                cell.font = Font(bold=True)
    if rows:
        last = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A{start_row}:{last}{start_row + len(rows)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    _autosize(ws, start_row)
    _print_setup(ws, start_row)


def _print_setup(ws: Worksheet, header_row: int) -> None:
    """Horizontal, ajustado al ancho de la página y con encabezado repetido."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{header_row}:{header_row}"


def _autosize(ws: Worksheet, start_row: int = 1) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            if cell.value is not None:
                length = max(len(line) for line in str(cell.value).splitlines() or [""])
                widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 9), MAX_WIDTH)


def write_excel(result: Result, path: Path, run_id: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    cats = categories(result)

    ws["A1"] = "Informe de conciliación de accesos"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("Fecha de corte", result.as_of.isoformat()),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Ejecución", run_id),
        ("Versión", f"hrrecon {__version__}"),
        ("Modo", "Solo lectura: generado a partir de archivos exportados"),
        *result.stats.items(),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(horizontal="left")
    row = len(meta) + 4
    ws.cell(row=row, column=1, value="Mensajes clave").font = Font(bold=True)
    for msg in key_messages(result):
        row += 1
        ws.cell(row=row, column=1, value=f"• {msg}")
    row += 2
    summary_rows = [
        [
            c.rule.code,
            c.rule.title,
            *[c.counts[s] for s in SEVERITIES],
            c.total,
            "Omitida: " + c.skipped if c.skipped else SHEET_NAMES[c.rule.code],
        ]
        for c in cats
    ]
    _table(
        ws,
        row,
        ["Regla", "Categoría", *[s.capitalize() for s in SEVERITIES], "Total", "Hoja"],
        summary_rows,
    )
    ws.freeze_panes = None
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 44

    for c in cats:
        sheet = wb.create_sheet(SHEET_NAMES[c.rule.code])
        sheet["A1"] = f"{c.rule.code} · {c.rule.title}"
        sheet["A1"].font = Font(bold=True, size=12)
        sheet["A2"] = c.rule.description
        sheet["A3"] = f"Acción recomendada: {c.rule.recommendation}"
        cols = COLUMNS[c.rule.code]
        if c.skipped:
            sheet["A5"] = f"Regla omitida: {c.skipped}"
            continue
        rows = [[cell_value(f, attr, None) for _, attr in cols] for f in c.findings]
        if not rows:
            sheet["A5"] = "Sin hallazgos."
        _table(sheet, 5, [h for h, _ in cols], rows)

    exc = wb.create_sheet("Excluidos")
    _table(
        exc,
        1,
        ["Tipo", "Identificador", "Nombre", "Motivo", "Reglas afectadas"],
        [[x.kind, x.identifier, x.name, x.reason, x.scope] for x in result.exclusions],
    )

    inputs = wb.create_sheet("Entradas")
    _table(
        inputs,
        1,
        ["Rol", "Archivo", "Filas", "Tamaño (bytes)", "SHA-256"],
        [[i.role, Path(i.path).name, i.rows, i.size, i.sha256] for i in result.inputs],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
