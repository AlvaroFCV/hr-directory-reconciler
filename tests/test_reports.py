from openpyxl import load_workbook

from hrrecon.report import render_html, write_excel
from hrrecon.report.common import SHEET_NAMES
from hrrecon.rules import RULES


def test_html_has_summary_and_every_category(sample):
    result, _, _ = sample
    html = render_html(result, "test-run")
    assert "Resumen ejecutivo" in html and "Exclusiones aplicadas" in html
    for rule in RULES:
        assert f'id="{rule.code}"' in html
        assert rule.title in html
    for i in result.inputs:
        assert i.sha256 in html
    assert "<script" not in html  # autocontenido, sin JavaScript
    assert "http://" not in html and "https://" not in html  # sin recursos externos


def test_html_escapes_values(sample):
    result, _, _ = sample
    result.findings[0].detail = "<script>alert(1)</script>"
    try:
        html = render_html(result, "x")
        assert "<script>alert(1)</script>" not in html
        assert "&lt;script&gt;" in html
    finally:
        result.findings[0].detail = ""


def test_excel_one_sheet_per_category(sample, tmp_path):
    result, _, _ = sample
    path = write_excel(result, tmp_path / "r.xlsx", "test-run")
    wb = load_workbook(path)
    expected = ["Resumen", *[SHEET_NAMES[r.code] for r in RULES], "Excluidos", "Entradas"]
    assert wb.sheetnames == expected
    for rule in RULES:
        ws = wb[SHEET_NAMES[rule.code]]
        data_rows = ws.max_row - 5  # título, descripción, acción, vacío, encabezado
        assert data_rows == len(result.by_rule(rule.code)), rule.code
    assert wb["Excluidos"].max_row - 1 == len(result.exclusions)
